from __future__ import annotations

from pathlib import Path
import json
import math
import shutil
import zipfile
import warnings

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy.integrate import solve_ivp
from scipy.special import voigt_profile
from scipy.stats import qmc, rankdata, norm, lognorm
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

SEED = 20260730
rng = np.random.default_rng(SEED)
ROOT = Path('/mnt/data/ros_rabi_literature_simulation_v1')
ROOT.mkdir(parents=True, exist_ok=True)

# Copy the extracted literature-parameter table created from the primary sources.
SRC_PARAM = Path('/mnt/data/ros_rabi_literature_simulation/literature_extracted_parameters.csv')
PARAM_CSV = ROOT / 'literature_extracted_parameters.csv'
shutil.copy2(SRC_PARAM, PARAM_CSV)
param_df = pd.read_csv(PARAM_CSV)

HC_EV_NM = 1239.841984
HBAR_EV_S = 6.582119569e-16
KB_EV_K = 8.617333262e-5
T_K = 310.15


def cm1_to_ev(x: float | np.ndarray) -> float | np.ndarray:
    return np.asarray(x) / 8065.544005


def inferred_lognormal_sigma(median: float, iqr_width: float) -> float:
    """Infer lognormal sigma from median and total IQR width under log-symmetric assumption."""
    z = 0.6744897501960817
    return float(np.arcsinh(iqr_width / (2.0 * median)) / z)


# -----------------------------------------------------------------------------
# Literature-constrained disease ROS distributions
# -----------------------------------------------------------------------------
ROS_SPECS = {
    'Healthy': (153.0, 112.0),
    'COVID-19': (1378.0, 1276.0),
    'Sepsis': (2120.0, 2024.0),
}
ROS_LOGN = {
    key: (math.log(median), inferred_lognormal_sigma(median, iqr))
    for key, (median, iqr) in ROS_SPECS.items()
}
ROS_FOLD_NOMINAL = {
    'Healthy': 1.0,
    'COVID-19': ROS_SPECS['COVID-19'][0] / ROS_SPECS['Healthy'][0],
    'Sepsis': ROS_SPECS['Sepsis'][0] / ROS_SPECS['Healthy'][0],
}

# -----------------------------------------------------------------------------
# Microscopic ROS-mediated two-state encounter: Rabi-like sensitivity
# The ET coupling and reorganization energy are EXTERNAL biomolecular priors.
# -----------------------------------------------------------------------------
V_CM1_NOM = 330.0
V_EV_NOM = float(cm1_to_ev(V_CM1_NOM))
OMEGA_RABI = 2.0 * V_EV_NOM / HBAR_EV_S
RABI_PERIOD_FS = 2.0 * np.pi / OMEGA_RABI * 1e15


def bloch_rhs(t, y, omega, delta, gamma2):
    u, v, w = y
    return [
        -delta * v - gamma2 * u,
        delta * u - omega * w - gamma2 * v,
        omega * v,
    ]


micro_t = np.linspace(0.0, 1000e-15, 1801)
MICRO_CASES = [
    {'label': 'Resonant upper bound', 'detuning_eV': 0.0, 'tau_phi_fs': 100.0},
    {'label': 'Moderate mismatch', 'detuning_eV': 0.10, 'tau_phi_fs': 100.0},
    {'label': 'DNA oxidation-gap prior', 'detuning_eV': 0.55, 'tau_phi_fs': 100.0},
    {'label': 'Protected/slow dephasing', 'detuning_eV': 0.10, 'tau_phi_fs': 1000.0},
    {'label': 'Fast dephasing', 'detuning_eV': 0.10, 'tau_phi_fs': 20.0},
]
microscopic_rows = []
microscopic_curves = {}
for case in MICRO_CASES:
    delta = case['detuning_eV'] / HBAR_EV_S
    gamma2 = 1.0 / (case['tau_phi_fs'] * 1e-15)
    sol = solve_ivp(
        bloch_rhs,
        (micro_t[0], micro_t[-1]),
        [0.0, 0.0, -1.0],
        t_eval=micro_t,
        args=(OMEGA_RABI, delta, gamma2),
        method='LSODA',
        rtol=1e-9,
        atol=1e-11,
    )
    pe = np.clip((1.0 + sol.y[2]) / 2.0, 0.0, 1.0)
    microscopic_curves[case['label']] = pe
    microscopic_rows.append({
        **case,
        'V_cm-1': V_CM1_NOM,
        'V_eV': V_EV_NOM,
        'Omega_Rabi_s-1': OMEGA_RABI,
        'Rabi_period_fs': RABI_PERIOD_FS,
        'maximum_excited_probability': float(np.max(pe)),
        'time_of_maximum_fs': float(micro_t[np.argmax(pe)] * 1e15),
    })
microscopic_df = pd.DataFrame(microscopic_rows)
microscopic_df.to_csv(ROOT / 'microscopic_Rabi_sensitivity.csv', index=False)

# Marcus ET rate surface from external DNA ET priors.
lambda_eV = float(cm1_to_ev(8070.0))
deltaG_eV = np.linspace(-2.0, 0.5, 501)
prefactor = (2.0 * np.pi / HBAR_EV_S) * V_EV_NOM**2 / np.sqrt(4.0 * np.pi * lambda_eV * KB_EV_K * T_K)
marcus_rate = prefactor * np.exp(-((deltaG_eV + lambda_eV) ** 2) / (4.0 * lambda_eV * KB_EV_K * T_K))
marcus_df = pd.DataFrame({'deltaG_eV': deltaG_eV, 'k_ET_s-1': marcus_rate})
marcus_df.to_csv(ROOT / 'Marcus_rate_sensitivity.csv', index=False)

# -----------------------------------------------------------------------------
# Spectral channels: centers/ranges/lifetimes from real measurements.
# Relative channel abundances and reaction orders remain model quantities.
# -----------------------------------------------------------------------------
CHANNELS = [
    {'name': 'Linoleic/triplet-carbonyl-like', 'center_nm': 475.0, 'lo_nm': 450.0, 'hi_nm': 490.0,
     'reaction_order': 1.0, 'prior_weight': 0.23, 'lifetime_us': 20.0, 'evidence': 'Tsuchida 2020 peak/range'},
    {'name': 'Protein/elastin-like', 'center_nm': 525.0, 'lo_nm': 500.0, 'hi_nm': 560.0,
     'reaction_order': 1.0, 'prior_weight': 0.13, 'lifetime_us': 20.0, 'evidence': 'Tsuchida 2020 peak/range'},
    {'name': 'Phospholipid-like', 'center_nm': 570.0, 'lo_nm': 530.0, 'hi_nm': 600.0,
     'reaction_order': 1.2, 'prior_weight': 0.25, 'lifetime_us': 20.0, 'evidence': 'Tsuchida 2020 peak/range'},
    {'name': 'Linolenic-like', 'center_nm': 575.0, 'lo_nm': 550.0, 'hi_nm': 620.0,
     'reaction_order': 1.5, 'prior_weight': 0.15, 'lifetime_us': 20.0, 'evidence': 'Tsuchida 2020 peak/range'},
    {'name': 'Singlet oxygen dimol 634 nm', 'center_nm': 634.0, 'lo_nm': 630.0, 'hi_nm': 638.0,
     'reaction_order': 2.0, 'prior_weight': 0.11, 'lifetime_us': 3.09, 'evidence': 'Lengfelder/Cadenas/Sies 1983'},
    {'name': 'Singlet oxygen dimol 703 nm', 'center_nm': 703.0, 'lo_nm': 698.0, 'hi_nm': 708.0,
     'reaction_order': 2.0, 'prior_weight': 0.08, 'lifetime_us': 3.09, 'evidence': 'Lengfelder/Cadenas/Sies 1983'},
    {'name': 'Singlet oxygen monomol 1270 nm', 'center_nm': 1270.0, 'lo_nm': 1260.0, 'hi_nm': 1280.0,
     'reaction_order': 2.0, 'prior_weight': 0.05, 'lifetime_us': 3.09, 'evidence': 'Lengfelder/Cadenas/Sies 1983'},
]
channel_df = pd.DataFrame(CHANNELS)
channel_df.to_csv(ROOT / 'spectral_channel_inputs.csv', index=False)


def sigma_from_relative_90(center, lo, hi):
    half_width = (abs(center - lo) + abs(hi - center)) / 2.0
    return max(half_width / np.sqrt(-2.0 * np.log(0.9)), 0.5)


for ch in CHANNELS:
    ch['sigma_nm'] = sigma_from_relative_90(ch['center_nm'], ch['lo_nm'], ch['hi_nm'])
    # The lifetime-derived natural linewidth is retained explicitly; measured band widths dominate.
    tau_s = ch['lifetime_us'] * 1e-6
    ch['natural_gamma_eV'] = HBAR_EV_S / (2.0 * tau_s)

WAVELENGTH_NM = np.linspace(300.0, 1350.0, 2101)


def voigt_in_energy_to_wavelength(lam_nm, center_nm, sigma_nm, natural_gamma_eV):
    energy = HC_EV_NM / lam_nm
    e0 = HC_EV_NM / center_nm
    sigma_e = HC_EV_NM * sigma_nm / center_nm**2
    profile_e = voigt_profile(energy - e0, sigma_e, natural_gamma_eV)
    # Jacobian |dE/dlambda|
    profile_l = profile_e * HC_EV_NM / lam_nm**2
    area = np.trapezoid(profile_l, lam_nm)
    return profile_l / max(area, 1e-300)


def gaussian_wavelength(lam_nm, center_nm, sigma_nm):
    x = np.exp(-0.5 * ((lam_nm - center_nm) / sigma_nm) ** 2)
    return x / max(np.trapezoid(x, lam_nm), 1e-300)


def lognormal_wavelength(lam_nm, center_nm, sigma_nm):
    shape = max(sigma_nm / center_nm, 0.01)
    # Approximate scale chosen to keep the mode near center_nm.
    scale = center_nm * np.exp(shape**2)
    x = lognorm.pdf(lam_nm, s=shape, scale=scale)
    return x / max(np.trapezoid(x, lam_nm), 1e-300)


BASIS_VOIGT = np.array([
    voigt_in_energy_to_wavelength(
        WAVELENGTH_NM,
        ch['center_nm'],
        ch['sigma_nm'],
        ch['natural_gamma_eV'],
    )
    for ch in CHANNELS
])
BASIS_GAUSSIAN = np.array([
    gaussian_wavelength(WAVELENGTH_NM, ch['center_nm'], ch['sigma_nm']) for ch in CHANNELS
])
BASIS_LOGNORMAL = np.array([
    lognormal_wavelength(WAVELENGTH_NM, ch['center_nm'], ch['sigma_nm']) for ch in CHANNELS
])

# Hamamatsu H7360-01 response, normalized because absolute sample geometry is not known.
DET_X = np.array([300.0, 400.0, 500.0, 600.0, 650.0])
DET_Y = np.array([1.4e5, 2.7e5, 1.7e5, 4.6e4, 0.0])


def detector_response(lam_nm, multipliers=None):
    y = DET_Y.copy()
    if multipliers is not None:
        y = y * np.asarray(multipliers)
    response = np.interp(lam_nm, DET_X, y, left=0.0, right=0.0)
    return response / max(response.max(), 1e-300)


DET_RESPONSE_NOM = detector_response(WAVELENGTH_NM)

BASE_ORDERS = np.array([ch['reaction_order'] for ch in CHANNELS], dtype=float)
BASE_WEIGHTS = np.array([ch['prior_weight'] for ch in CHANNELS], dtype=float)
BASE_WEIGHTS /= BASE_WEIGHTS.sum()
GAMMA_REL = np.array([1.0 / ch['lifetime_us'] for ch in CHANNELS], dtype=float)
MEAN_LAMBDA_V = np.array([
    np.trapezoid(WAVELENGTH_NM * b, WAVELENGTH_NM) for b in BASIS_VOIGT
])
BLUE_MASK = (WAVELENGTH_NM >= 400.0) & (WAVELENGTH_NM < 550.0)
RED_MASK = (WAVELENGTH_NM >= 600.0) & (WAVELENGTH_NM < 750.0)
BLUE_INT_V = np.array([
    np.trapezoid(b[BLUE_MASK], WAVELENGTH_NM[BLUE_MASK]) for b in BASIS_VOIGT
])
RED_INT_V = np.array([
    np.trapezoid(b[RED_MASK], WAVELENGTH_NM[RED_MASK]) for b in BASIS_VOIGT
])
DET_INT_V = np.array([
    np.trapezoid(b * DET_RESPONSE_NOM, WAVELENGTH_NM) for b in BASIS_VOIGT
])

# -----------------------------------------------------------------------------
# Generic ROS profiles. The clinical data constrain amplitude, not trajectory.
# -----------------------------------------------------------------------------
def normalized_pulse(t, rise, decay):
    shape = (1.0 - np.exp(-t / rise)) * np.exp(-t / decay)
    return shape / max(shape.max(), 1e-300)


def disease_ros_profile(t_h, fold, rise_h=5.0, decay_h=96.0):
    return fold * normalized_pulse(t_h, rise_h, decay_h)


def pma_ros_profile(t_s, amplitude=10.0, rise_s=360.0, decay_s=7200.0):
    return amplitude * normalized_pulse(t_s, rise_s, decay_s)

# -----------------------------------------------------------------------------
# Fast pathway integrator: Re -> intermediates -> chemical pump -> Pe -> UPE.
# -----------------------------------------------------------------------------
def run_pathway(
    t,
    rs,
    pars,
    weights=BASE_WEIGHTS,
    orders=BASE_ORDERS,
    basis=BASIS_VOIGT,
    detector_vector=DET_INT_V,
):
    n = len(t)
    m = len(CHANNELS)
    re = np.zeros(n, dtype=float)
    intermediates = np.zeros((m, n), dtype=float)

    # Heun integration for the slow oxidative/intermediate subsystem.
    for idx in range(1, n):
        dt = t[idx] - t[idx - 1]
        r0 = re[idx - 1]
        c0 = intermediates[:, idx - 1]

        source0 = pars['k_f'] * r0**orders / (
            pars['K_sat']**orders + r0**orders + 1e-300
        )
        dr0 = pars['k_se'] * rs[idx - 1] - pars['delta_re'] * r0
        dc0 = source0 - pars['delta_c'] * c0

        rp = max(0.0, r0 + dt * dr0)
        cp = np.maximum(0.0, c0 + dt * dc0)

        source1 = pars['k_f'] * rp**orders / (
            pars['K_sat']**orders + rp**orders + 1e-300
        )
        dr1 = pars['k_se'] * rs[idx] - pars['delta_re'] * rp
        dc1 = source1 - pars['delta_c'] * cp

        re[idx] = max(0.0, r0 + 0.5 * dt * (dr0 + dr1))
        intermediates[:, idx] = np.maximum(0.0, c0 + 0.5 * dt * (dc0 + dc1))

    # Xi = p/Gamma is the dimensionless chemical pump strength.
    xi = pars['gain'] * weights[:, None] * intermediates
    pe = xi / (1.0 + xi)
    channel_amplitudes = GAMMA_REL[:, None] * pe
    intrinsic_upe = channel_amplitudes.sum(axis=0)

    denom = channel_amplitudes.sum(axis=0)
    centroid = np.divide(
        MEAN_LAMBDA_V @ channel_amplitudes,
        denom,
        out=np.full(n, np.nan),
        where=denom > 0,
    )
    detector_upe = detector_vector @ channel_amplitudes
    blue = BLUE_INT_V @ channel_amplitudes
    red = RED_INT_V @ channel_amplitudes
    red_blue = np.divide(red, blue, out=np.full(n, np.nan), where=blue > 0)

    return {
        'time': t,
        'Rs': rs,
        'Re': re,
        'C': intermediates,
        'Xi': xi,
        'Pe': pe,
        'A': channel_amplitudes,
        'I': intrinsic_upe,
        'centroid': centroid,
        'detector_I': detector_upe,
        'red_blue': red_blue,
        'basis': basis,
    }

# -----------------------------------------------------------------------------
# Cell-level calibration to Burgos PMA/DPI ratio.
# Only composite gain/saturation/order shift are calibrated.
# -----------------------------------------------------------------------------
CELL_TIME_S = np.linspace(0.0, 9000.0, 301)
CELL_RATES = {
    'k_se': 0.0020,
    'delta_re': 0.00040,
    'k_f': 0.0015,
    'delta_c': 0.00050,
}
PMA_AMPLITUDE_NOM = 10.0  # normalized oxidative drive; no direct PMA ROS amplitude was reported.
PMA_RS = pma_ros_profile(CELL_TIME_S, amplitude=PMA_AMPLITUDE_NOM)
DPI_RESIDUAL_NOM = 0.50  # IC50-inspired nominal residual drive; tested over a broad range.
DPI_RS = DPI_RESIDUAL_NOM * PMA_RS

candidate_rows = []
for gain in np.logspace(-2.0, 1.5, 10):
    for ksat in np.logspace(0.5, 2.7, 12):
        for order_shift in (-0.2, 0.0, 0.2, 0.4, 0.6):
            orders = np.clip(BASE_ORDERS + order_shift, 0.7, 2.8)
            pars = {**CELL_RATES, 'gain': gain, 'K_sat': ksat}
            pma_out = run_pathway(CELL_TIME_S, PMA_RS, pars, orders=orders)
            dpi_out = run_pathway(CELL_TIME_S, DPI_RS, pars, orders=orders)
            ratio = dpi_out['detector_I'].max() / max(pma_out['detector_I'].max(), 1e-300)
            peak_occ = float(np.max(np.sum(pma_out['Pe'], axis=0)) / len(CHANNELS))
            score = abs(ratio - 0.45) + 0.05 * abs(peak_occ - 0.45)
            candidate_rows.append({
                'gain': gain,
                'K_sat': ksat,
                'order_shift': order_shift,
                'predicted_DPI_PMA_peak_ratio': ratio,
                'mean_peak_occupancy': peak_occ,
                'score': score,
            })
calibration_grid = pd.DataFrame(candidate_rows).sort_values('score').reset_index(drop=True)
calibration_grid.to_csv(ROOT / 'PMA_DPI_calibration_grid.csv', index=False)
best = calibration_grid.iloc[0]
NOMINAL_SHARED = {
    'gain': float(best['gain']),
    'K_sat': float(best['K_sat']),
    'order_shift': float(best['order_shift']),
}
NOMINAL_ORDERS = np.clip(BASE_ORDERS + NOMINAL_SHARED['order_shift'], 0.7, 2.8)
CELL_PARS = {**CELL_RATES, 'gain': NOMINAL_SHARED['gain'], 'K_sat': NOMINAL_SHARED['K_sat']}
PMA_OUT = run_pathway(CELL_TIME_S, PMA_RS, CELL_PARS, orders=NOMINAL_ORDERS)
DPI_OUT = run_pathway(CELL_TIME_S, DPI_RS, CELL_PARS, orders=NOMINAL_ORDERS)
NOMINAL_DPI_RATIO = float(DPI_OUT['detector_I'].max() / PMA_OUT['detector_I'].max())

# Inverse extraction: effective residual oxidative drive required for ratio 0.45.
residual_grid = np.linspace(0.10, 0.90, 161)
residual_results = []
for residual in residual_grid:
    out = run_pathway(CELL_TIME_S, residual * PMA_RS, CELL_PARS, orders=NOMINAL_ORDERS)
    ratio = out['detector_I'].max() / PMA_OUT['detector_I'].max()
    residual_results.append({'residual_oxidative_drive': residual, 'predicted_peak_ratio': ratio})
residual_df = pd.DataFrame(residual_results)
idx_best_res = (residual_df['predicted_peak_ratio'] - 0.45).abs().idxmin()
EFFECTIVE_DPI_RESIDUAL = float(residual_df.loc[idx_best_res, 'residual_oxidative_drive'])
residual_df.to_csv(ROOT / 'DPI_effective_residual_extraction.csv', index=False)

# -----------------------------------------------------------------------------
# Disease-scale nominal simulation: separate slow kinetics in h^-1.
# -----------------------------------------------------------------------------
DISEASE_TIME_H = np.linspace(0.0, 48.0, 241)
DISEASE_RATES = {
    'k_se': 0.25,
    'delta_re': 0.18,
    'k_f': 0.40,
    'delta_c': 0.35,
    'gain': NOMINAL_SHARED['gain'],
    'K_sat': NOMINAL_SHARED['K_sat'],
}
NOMINAL_DISEASE = {
    condition: run_pathway(
        DISEASE_TIME_H,
        disease_ros_profile(DISEASE_TIME_H, fold),
        DISEASE_RATES,
        orders=NOMINAL_ORDERS,
    )
    for condition, fold in ROS_FOLD_NOMINAL.items()
}

# -----------------------------------------------------------------------------
# LHS uncertainty propagation
# -----------------------------------------------------------------------------
N_LHS = 250
N_DIM = 15
U = qmc.LatinHypercube(d=N_DIM, seed=SEED).random(N_LHS)


def lin(u, a, b):
    return a + (b - a) * u


def log_factor(u, lo, hi):
    return 10 ** lin(u, np.log10(lo), np.log10(hi))


uncertainty_records = []
curve_store = {key: [] for key in ROS_FOLD_NOMINAL}
parameter_samples = []

for sample_idx, u in enumerate(U):
    # Disease ROS values sampled from inferred lognormal distributions.
    h = np.exp(ROS_LOGN['Healthy'][0] + ROS_LOGN['Healthy'][1] * norm.ppf(np.clip(u[0], 1e-6, 1 - 1e-6)))
    c = np.exp(ROS_LOGN['COVID-19'][0] + ROS_LOGN['COVID-19'][1] * norm.ppf(np.clip(u[1], 1e-6, 1 - 1e-6)))
    s = np.exp(ROS_LOGN['Sepsis'][0] + ROS_LOGN['Sepsis'][1] * norm.ppf(np.clip(u[2], 1e-6, 1 - 1e-6)))
    folds = {'Healthy': 1.0, 'COVID-19': c / h, 'Sepsis': s / h}

    pars = {
        'gain': NOMINAL_SHARED['gain'] * log_factor(u[3], 0.40, 2.50),
        'K_sat': NOMINAL_SHARED['K_sat'] * log_factor(u[4], 0.50, 2.00),
        'k_se': DISEASE_RATES['k_se'] * log_factor(u[5], 0.50, 2.00),
        'delta_re': DISEASE_RATES['delta_re'] * log_factor(u[6], 0.50, 2.00),
        'k_f': DISEASE_RATES['k_f'] * log_factor(u[7], 0.50, 2.00),
        'delta_c': DISEASE_RATES['delta_c'] * log_factor(u[8], 0.50, 2.00),
    }
    order_shift = NOMINAL_SHARED['order_shift'] + lin(u[9], -0.25, 0.25)
    orders = np.clip(BASE_ORDERS + order_shift, 0.7, 2.8)

    # Relative channel abundance is uncertain and is not extracted from cross-sample peak intensities.
    tilt = np.exp(lin(u[10], -0.40, 0.40) * np.linspace(-1.0, 1.0, len(CHANNELS)))
    weights = BASE_WEIGHTS * tilt
    weights /= weights.sum()

    rise_h = lin(u[11], 3.5, 7.5)
    decay_h = lin(u[12], 72.0, 144.0)

    det_multipliers = np.array([
        lin(u[13], 0.88, 1.12),
        lin(1.0 - u[13], 0.88, 1.12),
        lin(u[14], 0.88, 1.12),
        lin(1.0 - u[14], 0.88, 1.12),
        1.0,
    ])
    det_response = detector_response(WAVELENGTH_NM, det_multipliers)
    det_vector = np.array([
        np.trapezoid(b * det_response, WAVELENGTH_NM) for b in BASIS_VOIGT
    ])

    outputs = {}
    for condition, fold in folds.items():
        output = run_pathway(
            DISEASE_TIME_H,
            disease_ros_profile(DISEASE_TIME_H, fold, rise_h, decay_h),
            pars,
            weights=weights,
            orders=orders,
            detector_vector=det_vector,
        )
        outputs[condition] = output
        if sample_idx < 100:
            curve_store[condition].append(output['I'])

    # Cell consistency run, with cell rates and sampled shared pump mapping.
    pma_pars = {
        **CELL_RATES,
        'gain': pars['gain'],
        'K_sat': pars['K_sat'],
    }
    pma_amp = lin(u[11], 5.0, 20.0)
    pma_rs = pma_ros_profile(CELL_TIME_S, amplitude=pma_amp, rise_s=lin(u[12], 240, 540), decay_s=lin(u[13], 5400, 9000))
    dpi_residual = lin(u[14], 0.35, 0.65)
    pma = run_pathway(CELL_TIME_S, pma_rs, pma_pars, weights=weights, orders=orders, detector_vector=det_vector)
    dpi = run_pathway(CELL_TIME_S, dpi_residual * pma_rs, pma_pars, weights=weights, orders=orders, detector_vector=det_vector)
    dpi_ratio = dpi['detector_I'].max() / max(pma['detector_I'].max(), 1e-300)

    row = {
        'sample': sample_idx,
        'ROS_fold_COVID': folds['COVID-19'],
        'ROS_fold_Sepsis': folds['Sepsis'],
        'gain': pars['gain'],
        'K_sat': pars['K_sat'],
        'k_se': pars['k_se'],
        'delta_re': pars['delta_re'],
        'k_f': pars['k_f'],
        'delta_c': pars['delta_c'],
        'order_shift': order_shift,
        'rise_h': rise_h,
        'decay_h': decay_h,
        'dpi_residual_drive': dpi_residual,
        'DPI_PMA_peak_ratio': dpi_ratio,
    }

    healthy_absolute_threshold = 0.10 * outputs['Healthy']['I'].max()
    for condition, output in outputs.items():
        peak = output['I'].max()
        peak_idx = int(np.argmax(output['I']))
        relative_idx = np.where(output['I'] >= 0.10 * peak)[0]
        absolute_idx = np.where(output['I'] >= healthy_absolute_threshold)[0]
        row.update({
            f'{condition}_peak': peak,
            f'{condition}_AUC': np.trapezoid(output['I'], output['time']),
            f'{condition}_self_onset_h': output['time'][relative_idx[0]] if len(relative_idx) else np.nan,
            f'{condition}_healthy_threshold_onset_h': output['time'][absolute_idx[0]] if len(absolute_idx) else np.nan,
            f'{condition}_centroid_nm': output['centroid'][peak_idx],
            f'{condition}_red_blue': output['red_blue'][peak_idx],
        })

    uncertainty_records.append(row)
    parameter_samples.append(row)

uncertainty_df = pd.DataFrame(uncertainty_records)
uncertainty_df.to_csv(ROOT / 'uncertainty_draws.csv', index=False)

# Summary outputs
summary_rows = []
for condition in ROS_FOLD_NOMINAL:
    metrics = {
        'Peak intrinsic UPE': uncertainty_df[f'{condition}_peak'],
        'Cumulative intrinsic UPE': uncertainty_df[f'{condition}_AUC'],
        'Self-normalized onset (h)': uncertainty_df[f'{condition}_self_onset_h'],
        'Onset at 10% healthy peak (h)': uncertainty_df[f'{condition}_healthy_threshold_onset_h'],
        'Spectral centroid at peak (nm)': uncertainty_df[f'{condition}_centroid_nm'],
        'Red-to-blue ratio at peak': uncertainty_df[f'{condition}_red_blue'],
    }
    for metric_name, values in metrics.items():
        a = np.asarray(values, dtype=float)
        summary_rows.append({
            'Condition': condition,
            'Metric': metric_name,
            'Median': np.nanmedian(a),
            'P2.5': np.nanpercentile(a, 2.5),
            'P97.5': np.nanpercentile(a, 97.5),
            'IQR': np.nanpercentile(a, 75) - np.nanpercentile(a, 25),
        })
summary_df = pd.DataFrame(summary_rows)
summary_df.to_csv(ROOT / 'simulation_summary.csv', index=False)

fold_rows = []
for condition in ('COVID-19', 'Sepsis'):
    for metric in ('peak', 'AUC'):
        ratio = uncertainty_df[f'{condition}_{metric}'] / uncertainty_df[f'Healthy_{metric}']
        fold_rows.append({
            'Condition': condition,
            'Metric': f'{metric} fold vs healthy',
            'Median': np.median(ratio),
            'P2.5': np.percentile(ratio, 2.5),
            'P97.5': np.percentile(ratio, 97.5),
        })
fold_summary_df = pd.DataFrame(fold_rows)
fold_summary_df.to_csv(ROOT / 'fold_change_summary.csv', index=False)

# -----------------------------------------------------------------------------
# Structural spectral sensitivity: same channel amplitudes, alternative line shapes.
# -----------------------------------------------------------------------------
BASIS_MODELS = {
    'Voigt in energy': BASIS_VOIGT,
    'Gaussian in wavelength': BASIS_GAUSSIAN,
    'Log-normal in wavelength': BASIS_LOGNORMAL,
}
structural_rows = []
for model_name, basis in BASIS_MODELS.items():
    for condition, output in NOMINAL_DISEASE.items():
        peak_idx = int(np.argmax(output['I']))
        amplitude = output['A'][:, peak_idx]
        spectrum = amplitude @ basis
        area = np.trapezoid(spectrum, WAVELENGTH_NM)
        centroid = np.trapezoid(WAVELENGTH_NM * spectrum, WAVELENGTH_NM) / max(area, 1e-300)
        blue = np.trapezoid(spectrum[BLUE_MASK], WAVELENGTH_NM[BLUE_MASK])
        red = np.trapezoid(spectrum[RED_MASK], WAVELENGTH_NM[RED_MASK])
        structural_rows.append({
            'Line-shape model': model_name,
            'Condition': condition,
            'Centroid_nm': centroid,
            'Red_blue_ratio': red / max(blue, 1e-300),
        })
structural_df = pd.DataFrame(structural_rows)
structural_df.to_csv(ROOT / 'spectral_structural_sensitivity.csv', index=False)

# -----------------------------------------------------------------------------
# PRCC sensitivity for sepsis peak and centroid.
# -----------------------------------------------------------------------------
PRCC_COLUMNS = [
    'ROS_fold_Sepsis', 'gain', 'K_sat', 'k_se', 'delta_re', 'k_f',
    'delta_c', 'order_shift', 'rise_h', 'decay_h',
]


def calculate_prcc(df_x, y):
    x_rank = np.column_stack([rankdata(df_x[col]) for col in df_x.columns])
    y_rank = rankdata(y)
    rows = []
    for j, col in enumerate(df_x.columns):
        others = np.delete(x_rank, j, axis=1)
        design = np.column_stack([np.ones(len(y_rank)), others])
        bx = np.linalg.lstsq(design, x_rank[:, j], rcond=None)[0]
        by = np.linalg.lstsq(design, y_rank, rcond=None)[0]
        rx = x_rank[:, j] - design @ bx
        ry = y_rank - design @ by
        rows.append({'Parameter': col, 'PRCC': np.corrcoef(rx, ry)[0, 1]})
    return pd.DataFrame(rows)


prcc_peak_df = calculate_prcc(uncertainty_df[PRCC_COLUMNS], uncertainty_df['Sepsis_peak'])
prcc_peak_df['Output'] = 'Sepsis peak UPE'
prcc_centroid_df = calculate_prcc(uncertainty_df[PRCC_COLUMNS], uncertainty_df['Sepsis_centroid_nm'])
prcc_centroid_df['Output'] = 'Sepsis spectral centroid'
prcc_df = pd.concat([prcc_peak_df, prcc_centroid_df], ignore_index=True)
prcc_df.to_csv(ROOT / 'PRCC_results.csv', index=False)

# -----------------------------------------------------------------------------
# Compact Jansen/Saltelli-style Sobol estimates for sepsis peak and AUC.
# -----------------------------------------------------------------------------
SOBOL_PARAMS = ['ROS_fold_Sepsis', 'gain', 'K_sat', 'k_se', 'delta_re', 'k_f', 'delta_c', 'order_shift']
SOBOL_BOUNDS = {
    'ROS_fold_Sepsis': (7.0, 21.0),
    'gain': (NOMINAL_SHARED['gain'] * 0.4, NOMINAL_SHARED['gain'] * 2.5),
    'K_sat': (NOMINAL_SHARED['K_sat'] * 0.5, NOMINAL_SHARED['K_sat'] * 2.0),
    'k_se': (DISEASE_RATES['k_se'] * 0.5, DISEASE_RATES['k_se'] * 2.0),
    'delta_re': (DISEASE_RATES['delta_re'] * 0.5, DISEASE_RATES['delta_re'] * 2.0),
    'k_f': (DISEASE_RATES['k_f'] * 0.5, DISEASE_RATES['k_f'] * 2.0),
    'delta_c': (DISEASE_RATES['delta_c'] * 0.5, DISEASE_RATES['delta_c'] * 2.0),
    'order_shift': (NOMINAL_SHARED['order_shift'] - 0.25, NOMINAL_SHARED['order_shift'] + 0.25),
}
N_SOBOL = 128
rng_sobol = np.random.default_rng(SEED + 17)
A = np.column_stack([
    rng_sobol.uniform(*SOBOL_BOUNDS[p], size=N_SOBOL) for p in SOBOL_PARAMS
])
B = np.column_stack([
    rng_sobol.uniform(*SOBOL_BOUNDS[p], size=N_SOBOL) for p in SOBOL_PARAMS
])


def evaluate_sobol_matrix(matrix):
    peaks = np.empty(len(matrix))
    aucs = np.empty(len(matrix))
    for idx, row in enumerate(matrix):
        d = dict(zip(SOBOL_PARAMS, row))
        pars = {
            'gain': d['gain'], 'K_sat': d['K_sat'], 'k_se': d['k_se'],
            'delta_re': d['delta_re'], 'k_f': d['k_f'], 'delta_c': d['delta_c'],
        }
        orders = np.clip(BASE_ORDERS + d['order_shift'], 0.7, 2.8)
        out = run_pathway(
            DISEASE_TIME_H,
            disease_ros_profile(DISEASE_TIME_H, d['ROS_fold_Sepsis']),
            pars,
            orders=orders,
        )
        peaks[idx] = out['I'].max()
        aucs[idx] = np.trapezoid(out['I'], out['time'])
    return peaks, aucs


YA_peak, YA_auc = evaluate_sobol_matrix(A)
YB_peak, YB_auc = evaluate_sobol_matrix(B)
V_peak = np.var(np.concatenate([YA_peak, YB_peak]), ddof=1)
V_auc = np.var(np.concatenate([YA_auc, YB_auc]), ddof=1)
sobol_rows = []
for j, parameter in enumerate(SOBOL_PARAMS):
    AB = A.copy()
    AB[:, j] = B[:, j]
    YAB_peak, YAB_auc = evaluate_sobol_matrix(AB)
    s1_peak = np.mean(YB_peak * (YAB_peak - YA_peak)) / max(V_peak, 1e-300)
    st_peak = np.mean((YA_peak - YAB_peak) ** 2) / max(2.0 * V_peak, 1e-300)
    s1_auc = np.mean(YB_auc * (YAB_auc - YA_auc)) / max(V_auc, 1e-300)
    st_auc = np.mean((YA_auc - YAB_auc) ** 2) / max(2.0 * V_auc, 1e-300)
    sobol_rows.append({
        'Parameter': parameter,
        'S1_peak': s1_peak,
        'ST_peak': st_peak,
        'S1_AUC': s1_auc,
        'ST_AUC': st_auc,
    })
sobol_df = pd.DataFrame(sobol_rows)
sobol_df.to_csv(ROOT / 'Sobol_results.csv', index=False)

# -----------------------------------------------------------------------------
# Save nominal trajectories and spectra.
# -----------------------------------------------------------------------------
trajectory_rows = []
spectrum_rows = []
for condition, output in NOMINAL_DISEASE.items():
    for idx, t_value in enumerate(output['time']):
        trajectory_rows.append({
            'Condition': condition,
            'time_h': t_value,
            'signaling_ROS': output['Rs'][idx],
            'emission_relevant_oxidation': output['Re'][idx],
            'intrinsic_UPE': output['I'][idx],
            'detector_weighted_UPE': output['detector_I'][idx],
            'spectral_centroid_nm': output['centroid'][idx],
            'red_blue_ratio': output['red_blue'][idx],
        })
    peak_idx = int(np.argmax(output['I']))
    amplitude = output['A'][:, peak_idx]
    intrinsic_spectrum = amplitude @ BASIS_VOIGT
    detector_spectrum = intrinsic_spectrum * DET_RESPONSE_NOM
    for idx, lam in enumerate(WAVELENGTH_NM):
        spectrum_rows.append({
            'Condition': condition,
            'wavelength_nm': lam,
            'intrinsic_spectrum': intrinsic_spectrum[idx],
            'H7360_normalized_response': DET_RESPONSE_NOM[idx],
            'detector_weighted_spectrum': detector_spectrum[idx],
        })
trajectory_df = pd.DataFrame(trajectory_rows)
spectrum_df = pd.DataFrame(spectrum_rows)
trajectory_df.to_csv(ROOT / 'nominal_disease_trajectories.csv', index=False)
spectrum_df.to_csv(ROOT / 'nominal_spectra.csv', index=False)

# -----------------------------------------------------------------------------
# Figures
# -----------------------------------------------------------------------------
plt.figure(figsize=(8.2, 5.2))
for label, pe in microscopic_curves.items():
    plt.plot(micro_t * 1e15, pe, label=label)
plt.xlabel('Reactive-encounter time (fs)')
plt.ylabel('Excited-state probability')
plt.legend(fontsize=8)
plt.tight_layout()
plt.savefig(ROOT / 'Figure_1_microscopic_Rabi_like_dynamics.png', dpi=300)
plt.close()

plt.figure(figsize=(8.2, 5.2))
plt.semilogy(deltaG_eV, marcus_rate)
plt.axvline(-lambda_eV, linestyle='--', label=r'$\Delta G=-\lambda$')
plt.xlabel('Reaction free energy, ΔG (eV)')
plt.ylabel('Marcus electron-transfer rate (s$^{-1}$)')
plt.legend()
plt.tight_layout()
plt.savefig(ROOT / 'Figure_2_Marcus_rate_sensitivity.png', dpi=300)
plt.close()

plt.figure(figsize=(8.2, 5.2))
for condition, output in NOMINAL_DISEASE.items():
    normalized = output['I'] / max(NOMINAL_DISEASE['Healthy']['I'].max(), 1e-300)
    plt.plot(output['time'], normalized, label=condition)
plt.xlabel('Time (h)')
plt.ylabel('Intrinsic UPE relative to healthy peak')
plt.legend()
plt.tight_layout()
plt.savefig(ROOT / 'Figure_3_disease_UPE_predictions.png', dpi=300)
plt.close()

plt.figure(figsize=(8.2, 5.2))
plt.plot(CELL_TIME_S, PMA_OUT['detector_I'] / max(PMA_OUT['detector_I'].max(), 1e-300), label='PMA')
plt.plot(CELL_TIME_S, DPI_OUT['detector_I'] / max(PMA_OUT['detector_I'].max(), 1e-300), label='PMA + nominal DPI')
plt.axhspan(0.40, 0.50, alpha=0.15, label='Reported peak-ratio range')
plt.xlabel('Time after stimulation (s)')
plt.ylabel('H7360-band UPE normalized to PMA peak')
plt.legend()
plt.tight_layout()
plt.savefig(ROOT / 'Figure_4_PMA_DPI_pathway_calibration.png', dpi=300)
plt.close()

plt.figure(figsize=(8.2, 5.2))
for condition, output in NOMINAL_DISEASE.items():
    peak_idx = int(np.argmax(output['I']))
    spectrum = output['A'][:, peak_idx] @ BASIS_VOIGT
    spectrum /= max(np.trapezoid(spectrum, WAVELENGTH_NM), 1e-300)
    centroid = np.trapezoid(WAVELENGTH_NM * spectrum, WAVELENGTH_NM)
    plt.plot(WAVELENGTH_NM, spectrum, label=f'{condition} ({centroid:.0f} nm)')
plt.xlim(350, 1350)
plt.xlabel('Wavelength (nm)')
plt.ylabel('Normalized intrinsic spectral density')
plt.legend()
plt.tight_layout()
plt.savefig(ROOT / 'Figure_5_intrinsic_spectral_predictions.png', dpi=300)
plt.close()

plt.figure(figsize=(8.2, 5.2))
for condition, output in NOMINAL_DISEASE.items():
    peak_idx = int(np.argmax(output['I']))
    spectrum = (output['A'][:, peak_idx] @ BASIS_VOIGT) * DET_RESPONSE_NOM
    spectrum /= max(np.trapezoid(spectrum, WAVELENGTH_NM), 1e-300)
    plt.plot(WAVELENGTH_NM, spectrum, label=condition)
plt.xlim(300, 700)
plt.xlabel('Wavelength (nm)')
plt.ylabel('Normalized H7360-01-weighted spectrum')
plt.legend()
plt.tight_layout()
plt.savefig(ROOT / 'Figure_6_detector_weighted_spectra.png', dpi=300)
plt.close()

plt.figure(figsize=(8.2, 5.2))
for condition, curves in curve_store.items():
    arr = np.asarray(curves)
    med = np.median(arr, axis=0)
    lo = np.percentile(arr, 2.5, axis=0)
    hi = np.percentile(arr, 97.5, axis=0)
    plt.plot(DISEASE_TIME_H, med, label=condition)
    plt.fill_between(DISEASE_TIME_H, lo, hi, alpha=0.18)
plt.xlabel('Time (h)')
plt.ylabel('Intrinsic UPE (model units)')
plt.legend()
plt.tight_layout()
plt.savefig(ROOT / 'Figure_7_uncertainty_bands.png', dpi=300)
plt.close()

plt.figure(figsize=(8.2, 5.2))
for model_name, group in structural_df.groupby('Line-shape model'):
    values = [float(group[group['Condition'] == c]['Centroid_nm'].iloc[0]) for c in ('Healthy', 'COVID-19', 'Sepsis')]
    plt.plot(('Healthy', 'COVID-19', 'Sepsis'), values, marker='o', label=model_name)
plt.ylabel('Spectral centroid (nm)')
plt.legend(fontsize=8)
plt.tight_layout()
plt.savefig(ROOT / 'Figure_8_line_shape_sensitivity.png', dpi=300)
plt.close()

# -----------------------------------------------------------------------------
# Styled Excel workbook
# -----------------------------------------------------------------------------
NAVY = '17365D'
GREEN = '008000'
GRAY = '666666'
ORANGE = 'FCE4D6'
LIGHT_BLUE = 'D9EAF7'
THIN_GRAY = Side(style='thin', color='D9E1F2')


def add_dataframe_sheet(wb, title, df, source_columns=()):
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
        for col_name in source_columns:
            if col_name in df.columns:
                idx = list(df.columns).index(col_name) + 1
                row[idx - 1].font = Font(color=GREEN)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for idx, col in enumerate(df.columns, 1):
        max_len = max(len(str(col)), max((len(str(v)) for v in df[col].head(100)), default=0))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 42)
    return ws


wb = Workbook()
wb.remove(wb.active)

readme = wb.create_sheet('README')
readme['A1'] = 'Literature-Constrained ROS–Rabi–UPE Simulation'
readme['A1'].font = Font(size=16, bold=True, color=NAVY)
readme['A3'] = (
    'The workbook separates directly measured inputs, derived quantities, external biomolecular '
    'electron-transfer priors, and composite model parameters. The Rabi-like calculation describes '
    'one reactive molecular encounter. Disease-scale UPE is modeled as an incoherent ensemble chemical pump.'
)
readme['A5'] = (
    'Clinical ROS data constrain only disease oxidative amplitudes. No infection-specific UPE trajectory '
    'or spectrum was used as a fitting target. The PMA/DPI ratio is used as a pathway-level calibration/consistency test.'
)
readme['A7'] = (
    'Spectral centers and broad ranges are taken from measured oxidative biomolecule and singlet-oxygen spectra. '
    'Relative emitter abundances remain uncertain and were not inferred from cross-sample peak intensities.'
)
readme['A9'] = 'Orange cells indicate caution or proxy-prior status. Green text denotes imported literature values or URLs.'
readme.column_dimensions['A'].width = 120
for row in readme.iter_rows():
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')
readme.sheet_view.showGridLines = False

ws_params = add_dataframe_sheet(wb, 'Extracted Parameters', param_df, source_columns=('source', 'url'))
for row in ws_params.iter_rows(min_row=2):
    status = str(row[list(param_df.columns).index('status')].value).lower()
    if 'prior' in status or 'derived' in status:
        for cell in row:
            cell.fill = PatternFill('solid', fgColor=ORANGE)
    for col_name in ('value', 'lower', 'upper'):
        idx = list(param_df.columns).index(col_name) + 1
        row[idx - 1].font = Font(color=GREEN)

add_dataframe_sheet(wb, 'Microscopic Rabi', microscopic_df)
add_dataframe_sheet(wb, 'Marcus ET', marcus_df)
add_dataframe_sheet(wb, 'Spectral Channels', channel_df)
add_dataframe_sheet(wb, 'PMA-DPI Calibration', calibration_grid.head(100))
add_dataframe_sheet(wb, 'DPI Residual Extraction', residual_df)
add_dataframe_sheet(wb, 'Simulation Summary', summary_df)
add_dataframe_sheet(wb, 'Fold Changes', fold_summary_df)
add_dataframe_sheet(wb, 'Spectral Sensitivity', structural_df)
add_dataframe_sheet(wb, 'PRCC', prcc_df)
add_dataframe_sheet(wb, 'Sobol', sobol_df)

xlsx_path = ROOT / 'Literature_Constrained_Parameters_and_Results.xlsx'
wb.save(xlsx_path)

# -----------------------------------------------------------------------------
# Metadata, report, and package
# -----------------------------------------------------------------------------
nominal_summary = []
healthy_peak = NOMINAL_DISEASE['Healthy']['I'].max()
healthy_auc = np.trapezoid(NOMINAL_DISEASE['Healthy']['I'], DISEASE_TIME_H)
for condition, output in NOMINAL_DISEASE.items():
    peak_idx = int(np.argmax(output['I']))
    fixed_threshold = 0.10 * healthy_peak
    threshold_indices = np.where(output['I'] >= fixed_threshold)[0]
    nominal_summary.append({
        'Condition': condition,
        'ROS_fold': ROS_FOLD_NOMINAL[condition],
        'Peak_UPE': output['I'].max(),
        'Peak_fold_vs_healthy': output['I'].max() / healthy_peak,
        'AUC': np.trapezoid(output['I'], DISEASE_TIME_H),
        'AUC_fold_vs_healthy': np.trapezoid(output['I'], DISEASE_TIME_H) / healthy_auc,
        'Onset_at_10pct_healthy_peak_h': output['time'][threshold_indices[0]] if len(threshold_indices) else np.nan,
        'Centroid_at_peak_nm': output['centroid'][peak_idx],
        'Red_blue_at_peak': output['red_blue'][peak_idx],
    })
nominal_summary_df = pd.DataFrame(nominal_summary)
nominal_summary_df.to_csv(ROOT / 'nominal_summary.csv', index=False)

proportion_in_target = float(np.mean((uncertainty_df['DPI_PMA_peak_ratio'] >= 0.40) & (uncertainty_df['DPI_PMA_peak_ratio'] <= 0.50)))

metadata = {
    'random_seed': SEED,
    'direct_clinical_ROS_folds': ROS_FOLD_NOMINAL,
    'microscopic_external_prior': {
        'V_cm-1': V_CM1_NOM,
        'V_eV': V_EV_NOM,
        'Omega_Rabi_s-1': OMEGA_RABI,
        'Rabi_period_fs': RABI_PERIOD_FS,
        'lambda_cm-1': 8070.0,
        'lambda_eV': lambda_eV,
        'caution': 'DNA electron-transfer prior; not measured for infection-associated lipid/protein oxidation.',
    },
    'nominal_shared_pump_parameters': NOMINAL_SHARED,
    'nominal_DPI_PMA_peak_ratio_at_residual_0.5': NOMINAL_DPI_RATIO,
    'effective_DPI_residual_drive_for_ratio_0.45': EFFECTIVE_DPI_RESIDUAL,
    'uncertainty_samples': N_LHS,
    'fraction_uncertainty_draws_in_reported_DPI_ratio_interval': proportion_in_target,
    'sobol_base_sample_size': N_SOBOL,
}
with open(ROOT / 'run_metadata.json', 'w') as f:
    json.dump(metadata, f, indent=2)

# Manuscript-ready text kept deliberately cautious.
fold_lookup = fold_summary_df.set_index(['Condition', 'Metric'])
results_text = f"""
# Literature-constrained simulation results

The microscopic encounter calculation used an external biomolecular electron-transfer coupling of
{V_CM1_NOM:.0f} cm^-1 ({V_EV_NOM:.4f} eV), corresponding to a Rabi-like period of
{RABI_PERIOD_FS:.1f} fs. This value is an external DNA electron-transfer prior and is not interpreted
as an infection-specific lipid- or protein-oxidation constant. Increasing detuning and dephasing
strongly suppressed coherent excited-state transfer, supporting use of an incoherent ensemble pump
for disease-scale simulations.

Clinical oxidative amplitudes were constrained by median LIT values of 153, 1378, and 2120 for
healthy, severe COVID-19, and sepsis, respectively. These data generated nominal ROS folds of
1.00, {ROS_FOLD_NOMINAL['COVID-19']:.2f}, and {ROS_FOLD_NOMINAL['Sepsis']:.2f}.
No disease-specific UPE amplitude or spectrum was imposed.

The nominal disease simulation produced peak-UPE folds of
{nominal_summary_df.loc[nominal_summary_df.Condition=='COVID-19','Peak_fold_vs_healthy'].iloc[0]:.2f}
for COVID-19 and
{nominal_summary_df.loc[nominal_summary_df.Condition=='Sepsis','Peak_fold_vs_healthy'].iloc[0]:.2f}
for sepsis. Corresponding cumulative-emission folds were
{nominal_summary_df.loc[nominal_summary_df.Condition=='COVID-19','AUC_fold_vs_healthy'].iloc[0]:.2f}
and
{nominal_summary_df.loc[nominal_summary_df.Condition=='Sepsis','AUC_fold_vs_healthy'].iloc[0]:.2f}.
These values are model predictions conditioned on the extracted ROS data and composite pathway
parameters; they are not direct clinical UPE measurements.

The PMA/DPI calibration used the reported 0.40–0.50 residual UPE interval. With a nominal 50%
residual oxidative drive, the model produced a detector-band peak ratio of {NOMINAL_DPI_RATIO:.3f}.
Inverse pathway calibration required an effective residual oxidative drive of
{EFFECTIVE_DPI_RESIDUAL:.3f} to reproduce a ratio of 0.45. This is an effective model quantity,
not a directly measured ROS fraction. Across uncertainty draws, {proportion_in_target*100:.1f}% of
predicted ratios fell within the reported interval.

Spectral centers and widths were constrained using measured oxidative-biomolecule peaks and the
singlet-oxygen bands at 634, 703, and 1270 nm. Disease changed channel amplitudes through the
mechanistic reaction-order model; intrinsic transition centers were not shifted by disease. The
spectral conclusion was repeated with Voigt, Gaussian, and log-normal profiles and should be
interpreted as hypothesis-generating until infection-specific spectra are measured.
""".strip()
(ROOT / 'MANUSCRIPT_RESULTS_DRAFT.md').write_text(results_text)

readme_text = f"""
# ROS–Rabi–UPE Literature-Constrained Simulation

This package implements three linked but distinct scales:

1. a femtosecond microscopic two-state electron-transfer calculation;
2. a 9000-s PMA/DPI respiratory-burst pathway calibration;
3. a 48-h literature-constrained disease-scale ensemble chemical-pump simulation.

## Key caution

The 330 cm^-1 electronic coupling and 8070 cm^-1 reorganization energy are external DNA
charge-transfer priors. They establish a plausible quantum timescale but are not direct measurements
of ROS reactions with lipids or proteins during infection.

The disease-scale model is constrained by clinical ROS amplitudes, not by infection-specific UPE
data. Its photon-emission and spectral outputs are therefore prospective mechanistic predictions.

Nominal PMA/DPI peak ratio: {NOMINAL_DPI_RATIO:.3f}
Effective residual oxidative drive for ratio 0.45: {EFFECTIVE_DPI_RESIDUAL:.3f}
Rabi-like period from external coupling prior: {RABI_PERIOD_FS:.1f} fs
""".strip()
(ROOT / 'README.md').write_text(readme_text)

# Save executable source.
shutil.copy2(Path(__file__), ROOT / 'ros_rabi_upe_simulation.py')

zip_path = Path('/mnt/data/ROS_Rabi_UPE_Literature_Constrained_Simulation.zip')
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as archive:
    for file in ROOT.iterdir():
        archive.write(file, arcname=file.name)

print('Build complete')
print('Nominal shared parameters:', NOMINAL_SHARED)
print('Nominal DPI/PMA ratio:', NOMINAL_DPI_RATIO)
print('Effective residual drive:', EFFECTIVE_DPI_RESIDUAL)
print('Rabi period (fs):', RABI_PERIOD_FS)
print('\nNominal disease summary:')
print(nominal_summary_df.round(4).to_string(index=False))
print('\nUncertainty fold summary:')
print(fold_summary_df.round(4).to_string(index=False))
print('\nTop Sobol total-order indices for peak:')
print(sobol_df.sort_values('ST_peak', ascending=False).head(5)[['Parameter','ST_peak']].round(4).to_string(index=False))
print('\nPackage:', zip_path)
